import io
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.database import get_db
from app import models
from app.auth import get_current_user
from app.services import workflow as wf

router = APIRouter(prefix="/api", tags=["reports"])

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="1A2347")
HEADER_FONT = Font(color="FFFFFF", bold=True)


@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), user=Depends(get_current_user)):
    query = db.query(models.Contract)
    if user.role == "Контрагент":
        query = query.filter(models.Contract.contractor_id == user.contractor_id)
    elif user.role not in ("Директор", "Админ", "Юрист", "Бухгалтер"):
        query = query.filter(models.Contract.initiator_email == user.email)
    contracts = query.all()
    my_pending = wf.pending_for_role(db, user.role)

    return {
        "contractsTotal": len(contracts),
        "contractsPending": sum(1 for c in contracts if c.status == wf.STATUS_PENDING),
        "contractsApproved": sum(1 for c in contracts if c.status == wf.STATUS_APPROVED),
        "contractsRejected": sum(1 for c in contracts if c.status == wf.STATUS_REJECTED),
        "myPendingApprovals": len(my_pending),
    }


@router.get("/log/{entity_id}")
def get_log(entity_id: str, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if user.role == "Контрагент":
        contract = db.query(models.Contract).filter(models.Contract.id == entity_id).first()
        if not contract or contract.contractor_id != user.contractor_id:
            raise HTTPException(404, "Не найдено.")
    rows = db.query(models.LogEntry).filter(models.LogEntry.entity_id == entity_id) \
        .order_by(models.LogEntry.timestamp.asc()).all()
    return [{"date": r.timestamp, "user": r.user_email, "role": r.role,
             "action": r.action, "comment": r.comment} for r in rows]


@router.get("/export/contracts.xlsx")
def export_contracts(db: Session = Depends(get_db), user=Depends(get_current_user)):
    query = db.query(models.Contract)
    if user.role == "Контрагент":
        query = query.filter(models.Contract.contractor_id == user.contractor_id)
    elif user.role not in ("Директор", "Админ", "Юрист", "Бухгалтер"):
        query = query.filter(models.Contract.initiator_email == user.email)
    rows = query.order_by(models.Contract.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Договоры"
    headers = ["№", "Дата", "Инициатор", "Контрагент", "ИНН", "Юрлицо",
               "Предмет", "Номер договора", "Цена за единицу, ₽", "Статус",
               "Стандартный", "Действует до", "Комментарий"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"

    for c in rows:
        ws.append([
            c.id, c.created_at.strftime("%d.%m.%Y %H:%M") if c.created_at else "", c.initiator_fio,
            c.contractor_name, c.contractor_inn or "", c.legal_entity_name or "",
            c.subject, c.contract_number or "", float(c.price_per_unit or 0), c.status,
            ("Да" if c.is_standard else "Нет") if c.is_standard is not None else "—",
            c.valid_until.strftime("%d.%m.%Y") if c.valid_until else "", c.comment or "",
        ])
    for col in ws.columns:
        length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(length + 2, 10), 45)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type=XLSX_MIME,
        headers={"Content-Disposition": 'attachment; filename="contracts.xlsx"'},
    )
